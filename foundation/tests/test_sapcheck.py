"""
مقایسه با SAP: خواندن خروجی Excel/متنی SAP و تشخیص هر اختلاف.
(خروجی واقعی SAP هنوز در دست نیست؛ این آزمون‌ها سازوکار خواندن و مقایسه را می‌سنجند.)
"""
import json
from pathlib import Path

import pytest

from structural import sapcheck as S
from test_web import app  # noqa: F401

DATA = Path(__file__).parent / "data" / "sap"


@pytest.fixture(scope="module")
def la(tmp_path_factory):
    model, res, checks, combos = S.run(DATA / "LA.s2k")
    tables = S.program_tables(model, res, checks, combos)
    path = tmp_path_factory.mktemp("sap") / "LA_sap.xlsx"
    S.write_excel(tables, path, "LA")
    return tables, path


def test_program_tables_follow_sap_format(la):
    tables, _ = la
    assert {"Joint", "OutputCase", "F1", "M3"} <= set(tables["reactions"][0])
    cases = {r["OutputCase"] for r in tables["reactions"]}
    assert {"DEAD", "COMB1", "COMB8"} <= cases
    assert all(0 < r["Ratio"] < 5 for r in tables["steel"])


def test_same_numbers_pass(la):
    _, path = la
    rep = S.compare(DATA / "LA.s2k", path)
    assert rep["found"] == ["displacements", "reactions", "steel"]
    assert rep["ok"] and rep["compared"] > 1000


def test_changed_number_is_caught(la, tmp_path):
    from openpyxl import load_workbook
    _, path = la
    wb = load_workbook(path)
    ws = wb["Joint Reactions"]
    col = [c.value for c in ws[2]].index("F3") + 1
    ws.cell(row=4, column=col).value = ws.cell(row=4, column=col).value * 1.05 + 10
    ws2 = wb.worksheets[2]
    rc = [c.value for c in ws2[2]].index("Ratio") + 1
    ws2.cell(row=4, column=rc).value += 0.1
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    rep = S.compare(DATA / "LA.s2k", bad)
    assert not rep["ok"]
    assert sum(not r["ok"] for r in rep["reactions"]) == 1
    assert sum(not r["ok"] for r in rep["steel"]) == 1


def test_sap_text_export(tmp_path):
    """خروجی متنی SAP (همان قالب s2k، اعشار «/»)."""
    model, res, checks, combos = S.run(DATA / "PI.s2k")
    rows = S.program_tables(model, res, checks, combos)["reactions"][:4]
    lines = ['TABLE:  "JOINT REACTIONS"']
    for r in rows:
        vals = "   ".join(f"{k}={r[k]:.6f}".replace(".", "/") for k in S.REACTION)
        lines.append(f"   Joint={r['Joint']}   OutputCase={r['OutputCase']}   CaseType=LinStatic   {vals}")
    p = tmp_path / "out.s2k"
    p.write_text("\r\n".join(lines) + "\r\n", encoding="latin-1")
    got = S.read_sap_results(p)
    assert len(got["reactions"]) == 4
    rep = S.compare(DATA / "PI.s2k", p)
    assert rep["ok"] and rep["compared"] == 4


def test_verify_page(app, la):
    from test_web import _login
    c = app.app.test_client()
    h = {"X-CSRF-Token": _login(c)}
    _, path = la
    with open(DATA / "LA.s2k", "rb") as fm, open(path, "rb") as fr:
        r = c.post("/api/verify", data={"model": (fm, "LA.s2k"), "results": (fr, "LA_sap.xlsx")},
                   headers=h, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200 and d["ok"], json.dumps(d)[:300]
    assert c.get(d["excel"]).status_code == 200
    assert "مقایسه با SAP" in c.get("/verify").get_data(as_text=True)
