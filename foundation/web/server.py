"""
سامانه طراحی فونداسیون پایه تجهیزات پست — لایه وب.

اجرا:
    python server.py
سپس در مرورگر: http://127.0.0.1:5000
"""
import json, os, re, secrets, sys
from pathlib import Path
from dataclasses import asdict
from flask import (Flask, render_template, request, redirect, url_for, session,
                   flash, jsonify, send_file, abort, g)

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import database as db
import auth
import codeprofiles as cp
import settings as st
from calc_service import (build_equipment, run_layout, to_dict, layout_from_spec,
                          layout_problems, ALL_EQUIPMENT, VOLTAGE_LEVELS, EQUIPMENT_TYPES)
from config import ProjectConfig
from engine import from_config, GOVERNING_OPTIONS, BEARING_OPTIONS, RECOMMENDED, WHY
from equipment import CATALOG, outline_values  # noqa: F401
from seismic import FS_TABLE

VERSION = "2.5.0"      # در منوی کناری دیده می‌شود؛ نشانی فایل‌های css/js هم با آن عوض می‌شود

OUT = Path(os.environ.get("FOUNDATION_OUT") or Path(__file__).with_name("generated"))
OUT.mkdir(parents=True, exist_ok=True)


def _secret_key():
    """کلید نشست: از متغیر محیطی، وگرنه یک‌بار ساخته و کنار برنامه نگه داشته می‌شود."""
    env = os.environ.get("FOUNDATION_SECRET")
    if env:
        return env
    path = Path(__file__).with_name(".secret_key")
    if not path.exists():
        path.write_text(secrets.token_hex(32))
    return path.read_text().strip()


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 128 * 1024 * 1024       # نقشه جانمایی سه‌بعدی ده‌ها مگابایت است
app.secret_key = _secret_key()
app.jinja_env.add_extension("jinja2.ext.do")


@app.before_request
def _before():
    g._user = None
    auth.check_csrf()


@app.context_processor
def _inject():
    return {"user": auth.current_user(), "ROLES": auth.ROLES,
            "csrf_token": auth.csrf_token, "csrf_field": auth.csrf_field,
            "VERSION": VERSION}


# ================================================================= ورود
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = auth.verify(request.form.get("username", ""),
                           request.form.get("password", ""))
        if not user:
            flash("نام کاربری یا رمز عبور درست نیست.", "error")
        else:
            session.clear()
            session["uid"] = user["id"]
            g._user = user
            auth.record("ورود به سامانه")
            return redirect(auth.safe_next(request.args.get("next")) or url_for("dashboard"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    if auth.current_user():
        auth.record("خروج از سامانه")
    session.clear()
    return redirect(url_for("login"))


# ================================================================= داشبورد
@app.route("/")
@auth.login_required
def dashboard():
    stats = {
        "calculations": db.query("SELECT COUNT(*) c FROM calculations", one=True)["c"],
        "projects": db.query("SELECT COUNT(*) c FROM projects", one=True)["c"],
        "catalog": db.query("SELECT COUNT(*) c FROM equipment_catalog", one=True)["c"],
        "users": db.query("SELECT COUNT(*) c FROM users WHERE is_active=1", one=True)["c"],
    }
    recent = db.query(
        "SELECT c.*, u.full_name AS who FROM calculations c"
        " LEFT JOIN users u ON u.id=c.run_by ORDER BY c.run_at DESC LIMIT 8")
    return render_template("dashboard.html", stats=stats,
                           recent=[dict(r) for r in recent])


# ================================================================= محاسبه
PAGE_MODES = {
    "foundation": ("محاسبه فونداسیون", "طراحی و کنترل پی تجهیز: پایداری، ستون، آرماتور، میل مهار و نقشه ساخت"),
    "structure": ("سازه نگهدارنده", "تعریف، تحلیل و کنترل سازه فولادی تجهیز طبق AISC-ASD89 — جایگزین SAP"),
    "final": ("خروجی نهایی پست", "از نقشه جانمایی: طراحی همه پی‌ها و سازه‌ها، جمع مصالح و همه نقشه‌ها در یک فایل"),
}


@app.route("/calculate")
@auth.login_required
def calculate():
    return _calc_page("foundation")


@app.route("/structure")
@auth.login_required
def structure_page():
    return _calc_page("structure")


@app.route("/final")
@auth.login_required
def final_page():
    return _calc_page("final")


def _calc_page(mode):
    cat = db.query("SELECT * FROM equipment_catalog ORDER BY tag")
    subs = db.query("SELECT s.*, p.name AS project FROM substations s"
                    " LEFT JOIN projects p ON p.id=s.project_id ORDER BY s.code")
    return render_template("calculate.html",
                           catalog=[dict(r) for r in cat],
                           builtin={k: asdict(v) for k, v in ALL_EQUIPMENT.items()},
                           voltages=VOLTAGE_LEVELS, types=EQUIPMENT_TYPES,
                           guide=__import__("guide").by_id(),
                           # ترتیب فهرست تجهیز (tojson کلیدها را الفبایی می‌کند)
                           voltage_order={v: list(l["types"].items())
                                          for v, l in VOLTAGE_LEVELS.items()},
                           substations=[dict(r) for r in subs],
                           profiles=cp.listing(),
                           defaults=asdict(_office_config()),
                           angles=[a.name for a in __import__("structural.sections", fromlist=["ANGLES"]).ANGLES],
                           channels=[c.name for c in __import__("structural.sections", fromlist=["CHANNELS"]).CHANNELS],
                           options=design_options(), mode=mode,
                           page_title=PAGE_MODES[mode][0], page_sub=PAGE_MODES[mode][1])


def _office_config():
    """پیش‌فرض‌ها + تنظیمات دفتر (مثلاً قطر آرماتور)، برای مقدار اولیه فرم."""
    cfg = ProjectConfig()
    st.apply_to(cfg, st.load(), parts=st.ENGINEERING_PARTS)
    return cfg


def design_options():
    return {"governing": GOVERNING_OPTIONS, "bearing": BEARING_OPTIONS,
            "recommended": RECOMMENDED, "why": WHY}


class InputError(ValueError):
    pass


def _number(form, key, cast=float):
    val = form.get(key)
    if val is None or (isinstance(val, str) and not val.strip()):
        return None
    try:
        return cast(float(val))
    except (TypeError, ValueError):
        raise InputError(f"مقدار «{key}» عدد نیست: {val!r}")


def _choice(form, key, allowed, default):
    val = form.get(key)
    if val in (None, ""):
        return default
    if str(val) not in allowed:
        raise InputError(f"مقدار «{key}» باید یکی از {', '.join(allowed)} باشد")
    return str(val)


def config_from_form(form) -> ProjectConfig:
    """
    ساخت تنظیمات محاسبه از فرم.
    ترتیب: پیش‌فرض ← تنظیمات دفتر (آرماتور) ← مقادیر فرم.
    نتیجه کامل در اسنپ‌شات ذخیره می‌شود؛ تغییر بعدی تنظیمات روی این محاسبه اثر ندارد.
    """
    cfg = ProjectConfig()
    st.apply_to(cfg, st.load(), parts=st.ENGINEERING_PARTS)
    for group, keys in (("materials", ("fc", "fy", "cover", "lean")),
                        ("soil", ("q_base", "q_factor")),
                        ("wind", ("v_normal", "v_high")),
                        ("foundation", ("hp", "b", "tf", "min_projection", "L", "B")),
                        ("rebar", ("pad_dia", "col_dia", "tie_dia", "tie_spacing")),
                        ("steel", ("leg_width", "k_chord", "fy", "connection_factor", "panel"))):
        target = getattr(cfg, group)
        for key in keys:
            cur = getattr(target, key)
            val = _number(form, f"{group}.{key}", type(cur))
            if val is not None:
                setattr(target, key, val)

    for key in ("enabled", "feed_foundation"):
        val = form.get(f"steel.{key}")
        if val is not None:
            setattr(cfg.steel, key, val in (True, 1, "1", "true", "on"))
    from structural.sections import ANGLES, CHANNELS
    cfg.steel.mode = _choice(form, "steel.mode", ("design", "check"), cfg.steel.mode)
    cfg.steel.bracing = _choice(form, "steel.bracing", ("zigzag", "x"), cfg.steel.bracing)
    for key, allowed in (("chord", ANGLES), ("brace", ANGLES), ("strut", ANGLES), ("beam", CHANNELS)):
        val = form.get(f"steel.{key}")
        if val:
            names = [x.name for x in allowed]
            if val not in names:
                raise InputError(f"مقطع «{val}» در فهرست مقاطع نیست")
            setattr(cfg.steel, key, val)
    if not 0.2 <= cfg.steel.panel <= 2.0:
        raise InputError("ارتفاع پانل مهاربندی باید بین ۰٫۲ و ۲ متر باشد")
    if not 0.2 <= cfg.steel.leg_width <= 1.0:
        raise InputError("ضلع پایه مشبک باید بین ۰٫۲ و ۱ متر باشد")
    if cfg.steel.k_chord not in (1.0, 2.0):
        raise InputError("ضریب طول مؤثر نبشی اصلی باید ۱ یا ۲ باشد")
    fd = cfg.foundation
    if not (fd.b > 0 and fd.hp > 0 and fd.tf > 0):
        raise InputError("ضلع ستون b، ارتفاع ستون hp و ضخامت پی tf باید بزرگ‌تر از صفر باشند")

    s = cfg.seismic
    s.edition = int(_choice(form, "seismic.edition", ("4", "5"), str(s.edition)))
    for key in ("ss", "s1", "ie", "ru", "a", "b", "i", "r", "k_lateral"):
        val = _number(form, f"seismic.{key}")
        if val is not None:
            setattr(s, key, val)
    s.soil_class = _choice(form, "seismic.soil_class", tuple(FS_TABLE), s.soil_class)
    s.method = _choice(form, "seismic.method", ("rigid", "static"), s.method)
    d = cfg.design
    d.governing = _choice(form, "design.governing", tuple(GOVERNING_OPTIONS), d.governing)
    d.bearing = _choice(form, "design.bearing", tuple(BEARING_OPTIONS), d.bearing)
    return cfg


def _profile_for(form, edition):
    pid = form.get("code_profile_id")
    if pid:
        return int(pid)
    row = db.query("SELECT id FROM code_profiles WHERE edition=? AND is_active=1"
                   " ORDER BY id DESC LIMIT 1", (edition,), one=True)
    return row["id"] if row else None


@app.post("/api/calculate")
@auth.requires("engineer")
def api_calculate():
    form = request.get_json(silent=True)
    if not isinstance(form, dict):
        return jsonify({"error": "درخواست باید JSON باشد."}), 400
    try:
        cfg = config_from_form(form)
        eq = build_equipment(form.get("equipment") or {})
    except (TypeError, ValueError) as exc:
        return jsonify({"error": f"ورودی نامعتبر است: {exc}"}), 400

    # چیدمان پی: منفرد مربعی (۲۳۰/۴۰۰) یا مشترک مستطیلی با یک یا دو گروه تجهیز (۶۳)
    try:
        layout, layout_spec = layout_from_spec(form.get("layout"), eq)
    except (TypeError, ValueError, KeyError) as exc:
        return jsonify({"error": f"چیدمان پی نامعتبر است: {exc}"}), 400
    f = cfg.foundation
    if bool(f.L) != bool(f.B):
        return jsonify({"error": "برای کنترل یک پی معلوم، هر دو ضلع L و B را وارد کنید؛ "
                                 "یا هر دو را خالی بگذارید تا سامانه ابعاد را پیدا کند."}), 400
    if layout.square and f.L and abs(f.L - f.B) > 1e-9:
        return jsonify({"error": "پی منفرد مربعی است؛ L و B باید برابر باشند. برای پی "
                                 "مستطیلی نوع پی را «مشترک» انتخاب کنید."}), 400
    problems = layout_problems(layout, f.b, f.L, f.B)
    if problems:
        return jsonify({"error": " ".join(problems)}), 400

    # تجهیزاتی که سازه‌شان را سازنده می‌دهد: تا اعداد اوت‌لاین وارد نشود،
    # محاسبه اجرا نمی‌شود تا کسی سهواً با عدد نمونه نقشه نگیرد.
    base = ALL_EQUIPMENT.get(eq.tag)
    missing = []
    if base and base.requires_outline:
        raw = form.get("equipment") or {}
        LABELS = {"He": "ارتفاع تجهیز", "he": "مرکز ثقل تجهیز", "Ae": "سطح دید تجهیز",
                  "We": "وزن تجهیز", "Hs": "ارتفاع استراکچر", "As": "سطح دید استراکچر",
                  "Ws": "وزن استراکچر", "anchor_gauge": "گِیج میل مهار",
                  "base_plate": "ضلع صفحه کف", "conductor_points": "ارتفاع اتصال هادی",
                  "op_vertical": "بار قائم مانور", "op_horizontal": "بار افقی مانور"}
        for key in base.requires_outline:
            val = raw.get(key)
            if val in (None, "", 0) or (isinstance(val, str) and not val.strip()):
                missing.append(LABELS.get(key, key))
    if missing:
        return jsonify({"error": "این تجهیز سازه‌اش را سازنده می‌دهد؛ این مقادیر باید از "
                                 "اوت‌لاین سازنده وارد شوند: " + "، ".join(missing)}), 400

    res, seis, des, qty, bbs = run_layout(layout, cfg)
    if res is None:
        return jsonify({"error": "تا حد جست‌وجو ابعادی پیدا نشد که همه کنترل‌ها را پاس کند."}), 200

    payload = to_dict(res, seis, des, qty, bbs, eq, cfg)
    payload["id"] = _store(form, eq, cfg, layout_spec, payload)
    return jsonify(payload)


def _store(form, eq, cfg, layout_spec, payload):
    """ثبت محاسبه با اسنپ‌شات کامل ورودی (نقشه و گزارش فقط از همین ساخته می‌شوند)."""
    calc_id = db.execute(
        "INSERT INTO calculations (project_id, substation_id, equipment_tag,"
        " code_profile_id, inputs, results, status, run_by, run_at)"
        " VALUES (?,?,?,?,?,?,?,?,?)",
        (form.get("project_id") or None, form.get("substation_id") or None,
         eq.tag, _profile_for(form, cfg.seismic.edition),
         json.dumps({"equipment": asdict(eq), "config": asdict(cfg), "layout": layout_spec},
                    ensure_ascii=False),
         json.dumps(payload, ensure_ascii=False),
         "ok" if payload["ok"] else "ng",
         auth.current_user()["id"], db.now()))
    auth.record("اجرای محاسبه", "calculation", calc_id, {"tag": eq.tag})
    return calc_id


def _dxf_candidates(f, folder):
    """
    فایل بارگذاری‌شده ← فهرست DXFها: خود فایل، یا DXFهای داخل ZIP (پوشه نقشه‌ها)؛
    اول نام‌هایی که پلان/جانمایی‌اند، بعد بزرگ‌ترها.
    """
    import zipfile
    name = f.filename.lower()
    if name.endswith(".dxf"):
        path = folder / "drawing.dxf"
        f.save(path)
        return [path]
    if not name.endswith(".zip"):
        raise InputError("فایل باید DXF باشد، یا ZIP پوشه نقشه‌ها (در اتوکد: Save As ← DXF).")
    zpath = folder / "drawings.zip"
    f.save(zpath)
    out = []
    with zipfile.ZipFile(zpath) as z:
        members = [m for m in z.infolist() if m.filename.lower().endswith(".dxf") and not m.is_dir()]
        if sum(m.file_size for m in members) > 600 * 1024 * 1024:
            raise InputError("حجم نقشه‌های داخل ZIP بیش از حد است.")
        hint = lambda n: 0 if re.search(r"plan|layout|lyt|general", n, re.I) else 1
        members.sort(key=lambda m: (hint(Path(m.filename).name), -m.file_size))
        for i, m in enumerate(members):
            target = folder / f"{i:02d}.dxf"
            with z.open(m) as src, open(target, "wb") as dst:
                dst.write(src.read())
            out.append((target, Path(m.filename).name))
    if not out:
        raise InputError("در فایل ZIP هیچ نقشه DXF نیست (DWG را در اتوکد به DXF ذخیره کنید).")
    return out


def _layout_full(result, cfg, form, label):
    """
    طراحی کامل هر تیپ پی نقشه جانمایی و ثبت هر کدام به‌صورت یک محاسبه (گزارش و نقشه جدا)؛
    خلاصه کل پست: ابعاد، بتن، آرماتور، مقاطع سازه‌ها و فهرست کل فولاد.
    """
    import layoutplan as LP
    from structural.sections import CATALOG, STEEL
    types, steel_bill = [], {}
    tot = {"concrete": 0.0, "rebar": 0.0, "steel": 0.0, "pads": 0}
    for f, lay, c, run, ok in LP.design_all(result, cfg):
        res, seis, des, qty, bbs = run
        eq = lay.main
        payload = to_dict(res, seis, des, qty, bbs, eq, c)
        groups = []
        for i, g in enumerate(lay.groups):
            item = {"positions": [list(p) for p in g.positions]}
            if i:
                item["equipment"] = asdict(g.eq)
            groups.append(item)
        spec = {"kind": "combined", "source": f"نقشه جانمایی {label}: {f['name']}", "groups": groups}
        cid = _store(form, eq, c, spec, payload)
        n = f["count"]
        structs = []
        for st_ in payload.get("structures", []):
            secs = {g["group"]: g["section"] for g in st_["groups"]}
            structs.append({"tag": st_["tag"], "sections": secs, "weight": st_["weight_design"],
                            "ratio": st_["max_ratio"], "ok": st_["ok"]})
            tot["steel"] += st_["weight_design"] * n
            for g in st_["groups"]:
                sec = CATALOG.get(g["section"])
                row = steel_bill.setdefault(g["section"], {"section": g["section"], "length": 0.0,
                                                           "weight": 0.0, "count": 0})
                row["length"] += g["length"] * n
                row["count"] += g["count"] * n
                if sec is not None:
                    row["weight"] += g["length"] * n * sec.A * STEEL["gamma"]
        tot["concrete"] += qty["concrete"] * n
        tot["rebar"] += qty["rebar"] * n
        tot["pads"] += n
        types.append({"name": f["name"], "count": n, "L": c.foundation.L, "B": c.foundation.B,
                      "tf": c.foundation.tf, "b": c.foundation.b, "hp": c.foundation.hp,
                      "pedestals": sum(len(g.positions) for g in lay.groups),
                      "concrete": round(qty["concrete"], 2), "rebar": round(qty["rebar"], 1),
                      "ok": payload["ok"], "id": cid, "structures": structs,
                      "url": url_for("calculation_detail", cid=cid)})
    bill = sorted(steel_bill.values(), key=lambda r: (r["section"][0], CATALOG[r["section"]].A
                                                       if r["section"] in CATALOG else 0))
    for r in bill:
        r["length"], r["weight"] = round(r["length"], 1), round(r["weight"], 1)
    return {"types": types, "bill": bill,
            "totals": {k: round(v, 1) for k, v in tot.items()}}


def _outline_form():
    """عددهای تأییدشده اوت‌لاین که همراه فرم آمده: {نوع: {He, he, Ae, We}}."""
    import outline as OL
    try:
        return OL.overrides_ok(json.loads(request.form.get("outline") or "{}"))
    except (TypeError, ValueError, AttributeError) as exc:
        raise InputError(f"عددهای اوت‌لاین نامعتبر است: {exc}")


@app.post("/api/keyplan")
@auth.requires("engineer")
def api_keyplan():
    """
    کی‌پلن فونداسیون یا نقشه جانمایی (DXF یا ZIP پوشه نقشه‌ها).
      کی‌پلن  ← انواع پی، ابعاد، ستون‌ها و تجهیز هر ستون (keyplan.py)
      جانمایی ← جای تجهیزها و فضای موجود؛ همه پی‌ها با مشخصات ساختگاه فرم طراحی می‌شوند
                 (layoutplan.py) و کی‌پلن DXF ساخته می‌شود.
    فایل‌ها فقط خوانده و پاک می‌شوند.
    """
    import shutil
    import ezdxf
    from keyplan import read_foundations, detect_voltage as kp_voltage, plan as kp_plan
    import layoutplan as LP
    f = request.files.get("file")
    voltage = request.form.get("voltage") or "63"
    if not f or not f.filename:
        return jsonify({"error": "فایلی انتخاب نشده است."}), 400
    if voltage not in VOLTAGE_LEVELS:
        return jsonify({"error": "سطح ولتاژ نامعتبر است."}), 400
    folder = OUT / f"upload_{secrets.token_hex(8)}"
    folder.mkdir()
    try:
        cands = _dxf_candidates(f, folder)
        cands = [c if isinstance(c, tuple) else (c, f.filename) for c in cands]
        docs = []
        for path, label in cands:
            try:
                doc = ezdxf.readfile(str(path))
            except Exception:
                continue
            found = read_foundations(None, ALL_EQUIPMENT, voltage, doc=doc)
            if any(x.ok for x in found):
                # سطح ولتاژ از خود کی‌پلن: ستون‌های پی با کاتالوگ کدام سطح می‌خواند
                v_kp, found = kp_voltage(doc, ALL_EQUIPMENT, voltage)
                auth.record("خواندن کی‌پلن", "keyplan", None, {"file": f.filename, "types": len(found)})
                return jsonify({"kind": "keyplan", "file": label, "warning": "",
                                "voltage": v_kp, "voltage_changed": v_kp != voltage,
                                "foundations": [x.to_dict() for x in found],
                                "plan": {"stations": [],
                                         "pads": kp_plan(doc, {x.name for x in found if x.ok})}})
            items, axes = LP.read_items(None, doc=doc), LP.read_axes(doc)
            v = LP.detect_voltage(items) or voltage
            stations, _, _ = LP.site_stations(items, v, axes)
            n = sum(1 for s in stations if s.key)
            if n:
                docs.append((n, label, items, v, axes, doc))
            if n >= 3:
                break
        if not docs:
            return jsonify({"error": "نه بلاک پی کی‌پلن (مثل LA+CVT-3-2.5) پیدا شد، نه تجهیز در "
                                     "نقشه جانمایی (بلاک‌هایی با نام LA، CT، CB، DS، CVT، PI)."}), 400
        _, label, items, voltage_found, axes, doc = max(docs, key=lambda d: d[0])
        del docs
        try:
            form = json.loads(request.form.get("form") or "{}")
        except ValueError:
            form = {}
        if not form.get("soil.q_base") or not (form.get("seismic.a") or form.get("seismic.ss")):
            raise InputError("این فایل نقشه جانمایی است. برای طراحی پی‌ها ابتدا مشخصات ساختگاه "
                             "(زلزله، خاک و باد) را در فرم وارد کنید و سپس دوباره بارگذاری کنید.")
        cfg = config_from_form(form)
        gap = _number(form, "layout.gap")
        gap = 0.20 if gap is None else gap
        if not 0 <= gap <= 1:
            raise InputError("فاصله آزاد بین پی‌ها باید بین ۰ و ۱ متر باشد.")
        # محدوده FUTURE PLAN: از ابرهای همین نقشه یا برگه‌هایی از ZIP که همین فایل را xref کرده‌اند
        future = []
        if form.get("layout.future") != "design":
            future = LP.future_zones(doc)
            target = Path(label).stem
            for path, lab in cands:
                if lab == label or Path(path).stat().st_size > 30e6:
                    continue
                try:
                    future += LP.future_zones(ezdxf.readfile(str(path)), target)
                except Exception:
                    continue
        del doc
        with outline_values(voltage_found, _outline_form()):
            result = LP.design_layout(None, cfg, voltage_found, items=items, axes=axes, gap=gap,
                                      future=future)
            dxf = f"keyplan_from_layout_{secrets.token_hex(4)}.dxf"
            LP.plan_dxf(result, str(OUT / dxf))
            full = _layout_full(result, cfg, form, label) if request.form.get("full") == "1" else None
    except InputError as exc:
        return jsonify({"error": str(exc)}), 400
    except ValueError as exc:
        return jsonify({"error": f"نقشه خوانده نشد: {exc}"}), 400
    except Exception as exc:                      # هر خطای دیگر: پیام روشن، نه صفحه خطا
        app.logger.exception("layout/keyplan")
        return jsonify({"error": f"نقشه خوانده نشد ({type(exc).__name__}: {exc})"}), 500
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    auth.record("طراحی از نقشه جانمایی", "layout", None,
                {"file": f.filename, "types": len(result["foundations"])})
    return jsonify({"kind": "layout", "file": label, "voltage": voltage_found,
                    "voltage_changed": voltage_found != voltage,
                    "foundations": result["foundations"],
                    "plan": result["plan"], "notes": result["notes"], "unknown": result["unknown"],
                    "rows": [{"axis": r["axis"], "units": r["units"]} for r in result["rows"]],
                    "dxf": url_for("download", name=dxf), "full": full})


# ================================================================= اوت‌لاین سازنده
@app.post("/api/outline")
@auth.requires("engineer")
def api_outline():
    """
    نقشه اوت‌لاین سازنده (PDF، یک یا چند صفحه): نوع تجهیز و عددهای بار هر صفحه، کنار عدد
    کاتالوگ همان نوع. فقط پیشنهاد است؛ مهندس با دیدن تصویر صفحه تأیید می‌کند.
    """
    import outline as OL
    f = request.files.get("file")
    voltage = request.form.get("voltage") or "230"
    if not f or not f.filename.lower().endswith(".pdf"):
        return jsonify({"error": "نقشه اوت‌لاین را به‌صورت PDF بدهید."}), 400
    if voltage not in VOLTAGE_LEVELS:
        return jsonify({"error": "سطح ولتاژ نامعتبر است."}), 400
    oid = secrets.token_hex(8)
    folder = OUT / f"outline_{oid}"
    folder.mkdir()
    try:
        f.save(folder / "source.pdf")
        pages = OL.read_pdf(str(folder / "source.pdf"))
    except Exception as exc:
        app.logger.exception("outline")
        return jsonify({"error": f"PDF خوانده نشد ({type(exc).__name__}: {exc})"}), 400
    types = VOLTAGE_LEVELS[voltage]["types"]
    for p in pages:
        key = types.get(p["type"])
        eq = ALL_EQUIPMENT.get(key) if key else None
        p["catalog"] = {k: getattr(eq, k) for k in ("He", "he", "Ae", "We")} if eq else None
        p["catalog_key"] = key or ""
        p["image"] = url_for("outline_page", oid=oid, page=p["page"])
    (folder / "pages.json").write_text(json.dumps(pages, ensure_ascii=False))
    auth.record("خواندن اوت‌لاین", "outline", None,
                {"file": f.filename, "pages": len(pages),
                 "types": [p["type"] for p in pages]})
    scanned = sum(p["method"] == "none" for p in pages)
    return jsonify({"file": f.filename, "pages": pages, "ocr": OL.ocr_available(),
                    "types": {t: EQUIPMENT_TYPES.get(t, t) for t in types},
                    "warning": (f"{scanned} صفحه اسکن است و خواندن متن تصویر (OCR) روی این رایانه "
                                "نصب نیست؛ عددها را از تصویر همان صفحه وارد کنید."
                                if scanned else "")})


@app.route("/api/outline/<oid>/<int:page>.png")
@auth.login_required
def outline_page(oid, page):
    import outline as OL
    if not re.fullmatch(r"[0-9a-f]{16}", oid):
        abort(404)
    folder = OUT / f"outline_{oid}"
    src = folder / "source.pdf"
    if not src.is_file():
        abort(404)
    pages = {p["page"]: p for p in json.loads((folder / "pages.json").read_text())}
    if page not in pages:
        abort(404)
    png = folder / f"p{page}.png"
    if not png.is_file():
        png.write_bytes(OL.render_page(str(src), page, pages[page]["fields"]))
    return send_file(png, mimetype="image/png")


# ================================================================= تنظیمات نقشه
TITLE_LABELS = [
    ("DOCUMENT_TITLE", "عنوان نقشه"), ("DESC.1", "شرح ۱"), ("DESC.2", "شرح ۲"),
    ("PANEL_NAME", "نام پست"), ("DWG.", "شماره نقشه"), ("NO.P", "شماره صفحه"),
    ("T.P", "تعداد صفحات"), ("COUNT.", "شماره مدرک"), ("CONTRACT_NO.", "شماره قرارداد"),
    ("CONTRACTOR_DOC_CODE", "کد مدرک پیمانکار"),
    ("PREP", "تهیه‌کننده"), ("CHCK", "کنترل‌کننده"), ("APP", "تأییدکننده"), ("AUTH", "تصویب"),
    ("REV1", "بازنگری"), ("DATE1", "تاریخ بازنگری"), ("DESC1", "شرح بازنگری"),
    ("MOD.1", "اصلاح"), ("CHCK.1", "کنترل"), ("APPRD.1", "تأیید"),
]


@app.route("/settings", methods=["GET", "POST"])
@auth.login_required
def drawing_settings():
    if request.method == "POST":
        if auth.RANK[auth.current_user()["role"]] < auth.RANK["engineer"]:
            abort(403)
        data = {
            "notes": request.form.get("notes", ""),
            "title_block": {k: request.form.get("tb." + k, "").strip()
                            for k, _ in TITLE_LABELS},
            "drawing": {k: request.form.get("dw." + k, "").strip()
                        for k in ("scale", "detail_scale")},
            "rebar": {k: request.form.get("rb." + k, "").strip()
                      for k in ("pad_dia", "pad_spacing_max", "col_dia", "col_min_bars",
                                "tie_dia", "tie_spacing", "standee_dia")},
        }
        st.save(data, auth.current_user()["id"])
        for old in OUT.glob("sheet_*.svg"):          # جدول عنوان و یادداشت‌ها روی شیت اثر دارند
            old.unlink(missing_ok=True)
        auth.record("ویرایش تنظیمات نقشه", "settings")
        flash("تنظیمات نقشه ذخیره شد. از این پس روی نقشه‌های تازه اعمال می‌شود.", "ok")
        return redirect(url_for("drawing_settings"))

    data = st.load()
    return render_template("settings.html", data=data, labels=TITLE_LABELS,
                           defaults=asdict(ProjectConfig()))


# ================================================================= نقشه و گزارش
def _same_result(stored, fresh):
    """آیا بازتولید از اسنپ‌شات همان نتیجه ذخیره‌شده را داد؟"""
    a, b = stored["geometry"], fresh["geometry"]
    if any(abs(a[k] - b[k]) > 1e-9 for k in ("L", "B", "tf", "hp", "b")):
        return False
    for key in ("pad", "pedestal"):
        x, y = stored["design"][key], fresh["design"][key]
        if (x["bars"], x["dia"], x["spacing"]) != (y["bars"], y["dia"], y["spacing"]):
            return False
    return True


def _rebuild(cid):
    """
    بازتولید یک محاسبه ثبت‌شده از اسنپ‌شات (برای نقشه و پیش‌نمایش نقشه در صفحه).
    فقط تنظیمات ظاهری (جدول عنوان، مقیاس، یادداشت‌ها) از تنظیمات فعلی می‌آید؛ اگر
    بازتولید با نتیجه ذخیره‌شده نخواند، خطا برمی‌گردد.
    خروجی: ((res, seis, des, qty, bbs, eq, cfg)، None) یا (None، پاسخ خطا)
    """
    row = db.query("SELECT * FROM calculations WHERE id=?", (cid,), one=True)
    if not row:
        return None, (jsonify({"error": "چنین محاسبه‌ای ثبت نشده."}), 404)

    saved = json.loads(row["inputs"])
    cfg = ProjectConfig._from_dict(saved["config"])
    root = Path(__file__).resolve().parent.parent
    cfg.drawing.frame_file = str(root / Path(cfg.drawing.frame_file).name)
    cfg.drawing.template_file = str(root / Path(cfg.drawing.template_file).name)
    cfg.drawing.notes_file = str(root / Path(cfg.drawing.notes_file).name)

    conf = st.load()
    st.apply_to(cfg, conf, parts=st.PRESENTATION_PARTS)
    notes_text = (conf.get("notes") or "").strip()
    if notes_text:
        notes_file = OUT / f"notes_{cid}.txt"
        notes_file.write_text(notes_text, encoding="utf-8")
        cfg.drawing.notes_file = str(notes_file)

    eq = build_equipment(saved["equipment"])
    layout, _ = layout_from_spec(saved.get("layout"), eq)
    res, seis, des, qty, bbs = run_layout(layout, cfg)
    if res is None or not _same_result(json.loads(row["results"]),
                                       to_dict(res, seis, des, qty, bbs, eq, cfg)):
        return None, (jsonify({"error": "بازتولید محاسبه از اسنپ‌شات با نتیجه ذخیره‌شده نخواند؛ "
                                        "نقشه ساخته نشد. محاسبه را دوباره اجرا کنید."}), 409)
    return (res, seis, des, qty, bbs, eq, cfg), None


@app.get("/api/sheet/<int:cid>.svg")
@auth.login_required
def api_sheet(cid):
    """
    پیش‌نمایش نقشه ساخت در صفحه: همان شیت دوبعدی اتوکد (همان کد و همان اسنپ‌شات)
    به SVG — آنچه در صفحه دیده می‌شود دقیقاً همان است که در فایل DXF می‌رود.
    """
    cache = OUT / f"sheet_{cid}.svg"
    if cache.is_file():
        return send_file(cache, mimetype="image/svg+xml", max_age=0)
    built, err = _rebuild(cid)
    if err:
        return err
    res, seis, des, qty, bbs, eq, cfg = built
    try:
        from sheet_preview import sheet_svg
        cache.write_text(sheet_svg(res, eq, from_config(cfg)[0], qty, bbs, seis, des, cfg),
                         encoding="utf-8")
    except ImportError:
        return jsonify({"error": "کتابخانه ezdxf نصب نیست"}), 500
    except Exception as exc:
        app.logger.exception("sheet preview failed")
        return jsonify({"error": f"پیش‌نمایش نقشه ساخته نشد: {exc}"}), 500
    return send_file(cache, mimetype="image/svg+xml", max_age=0)


@app.post("/api/drawing/<int:cid>")
@auth.requires("engineer")
def api_drawing(cid):
    """
    تولید DXF از اسنپ‌شات یک محاسبه ذخیره‌شده.
    فقط تنظیمات ظاهری (جدول عنوان، مقیاس، یادداشت‌ها) از تنظیمات فعلی می‌آید؛
    هر چه روی عدد اثر دارد از خود اسنپ‌شات است، و اگر بازتولید با نتیجه
    ذخیره‌شده نخواند، نقشه ساخته نمی‌شود.
    """
    built, err = _rebuild(cid)
    if err:
        return err
    res, seis, des, qty, bbs, eq, cfg = built

    try:
        from outputs import make_outputs
        out = make_outputs(res, eq, from_config(cfg)[0], qty, bbs, seis, des, cfg,
                           str(OUT), f"{eq.tag}_{cid}")
    except ImportError:
        return jsonify({"error": "کتابخانه ezdxf نصب نیست: python -m pip install ezdxf"}), 500
    except Exception as exc:
        app.logger.exception("drawing failed")
        return jsonify({"error": f"نقشه ساخته نشد: {exc}"}), 500

    n2, n3 = Path(out["2d"]).name, Path(out["3d"]).name
    sap = [Path(x).name for x in out.get("sap", [])]
    db.execute("UPDATE calculations SET dxf_path=?, dxf3d_path=? WHERE id=?", (n2, n3, cid))
    auth.record("تولید نقشه", "calculation", cid, {"2d": n2, "3d": n3})
    return jsonify({"files": [
        {"kind": "2d", "name": n2, "url": url_for("download", name=n2),
         "label": f"نقشه دوبعدی (مقیاس ۱:{out['scale']:.0f})"},
        {"kind": "3d", "name": n3, "url": url_for("download", name=n3),
         "label": "مدل سه‌بعدی"}] + [
        {"kind": "sap", "name": n, "url": url_for("download", name=n),
         "label": "مدل SAP سازه"} for n in sap],
        "warnings": out["warnings"]})


@app.route("/files/<path:name>")
@auth.login_required
def download(name):
    target = (OUT / name).resolve()
    if OUT.resolve() not in target.parents or not target.is_file():
        abort(404)
    return send_file(target, as_attachment=True)


@app.post("/api/final/zip")
@auth.requires("engineer")
def api_final_zip():
    """
    خروجی نهایی پست در یک فایل: نقشه دوبعدی، مدل سه‌بعدی و مدل SAP هر تیپ پی، کی‌پلن ساخته‌شده از
    نقشه جانمایی و جدول خلاصه (Excel). ورودی: {"ids": [...], "summary": {...}, "keyplan": نام فایل}
    """
    import zipfile
    body = request.get_json(silent=True) or {}
    ids = [int(i) for i in body.get("ids", []) if str(i).isdigit()][:60]
    if not ids:
        return jsonify({"error": "هیچ محاسبه‌ای برای خروجی نیست."}), 400
    from outputs import make_outputs
    name = f"final_{secrets.token_hex(4)}.zip"
    warnings = []
    with zipfile.ZipFile(OUT / name, "w", zipfile.ZIP_DEFLATED) as z:
        for cid in ids:
            built, err = _rebuild(cid)
            if err:
                warnings.append(f"محاسبه {cid} بازتولید نشد")
                continue
            res, seis, des, qty, bbs, eq, cfg = built
            try:
                out = make_outputs(res, eq, from_config(cfg)[0], qty, bbs, seis, des, cfg,
                                   str(OUT), f"{eq.tag}_{cid}")
            except Exception as exc:
                app.logger.exception("final zip")
                warnings.append(f"محاسبه {cid}: {exc}")
                continue
            folder = body.get("names", {}).get(str(cid)) or f"{eq.tag}_{cid}"
            folder = re.sub(r"[^\w.+-]", "_", folder)
            for key in ("2d", "3d"):
                z.write(out[key], f"{folder}/{Path(out[key]).name}")
            for sp in out.get("sap", []):
                z.write(sp, f"{folder}/{Path(sp).name}")
        kp = body.get("keyplan")
        if kp:
            kp_path = (OUT / Path(kp).name).resolve()
            if OUT.resolve() in kp_path.parents and kp_path.is_file():
                z.write(kp_path, "KEYPLAN_from_layout.dxf")
        summ = body.get("summary")
        if summ:
            try:
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Foundations"
                ws.append(["تیپ پی", "تعداد", "L", "B", "tf", "ستون", "بتن هر پی m3", "آرماتور kg",
                           "سازه", "نبشی اصلی", "مهاربند", "افقی", "تیر سر", "وزن سازه kg", "نسبت تنش"])
                for t in summ.get("types", []):
                    sts = t.get("structures") or [{}]
                    for st_ in sts:
                        sec = st_.get("sections", {})
                        ws.append([t["name"], t["count"], t["L"], t["B"], t["tf"],
                                   f'{t["pedestals"]}x{t["b"]}', t["concrete"], t["rebar"],
                                   st_.get("tag", "سازنده"), sec.get("chord"), sec.get("brace"),
                                   sec.get("strut"), sec.get("beam"), st_.get("weight"), st_.get("ratio")])
                ws2 = wb.create_sheet("Steel")
                ws2.append(["مقطع", "تعداد عضو", "طول کل m", "وزن kg"])
                for r in summ.get("bill", []):
                    ws2.append([r["section"], r["count"], r["length"], r["weight"]])
                ws3 = wb.create_sheet("Totals")
                for k, lab in (("pads", "تعداد پی"), ("concrete", "بتن m3"), ("rebar", "آرماتور kg"),
                               ("steel", "فولاد سازه‌ها kg")):
                    ws3.append([lab, summ.get("totals", {}).get(k)])
                tmp = OUT / f"summary_{secrets.token_hex(3)}.xlsx"
                wb.save(tmp)
                z.write(tmp, "SUMMARY.xlsx")
                tmp.unlink(missing_ok=True)
            except ImportError:
                warnings.append("جدول خلاصه ساخته نشد: openpyxl نصب نیست")
    auth.record("خروجی نهایی", "final", None, {"count": len(ids)})
    return jsonify({"url": url_for("download", name=name), "warnings": warnings})


# ================================================================= تاریخچه
@app.route("/history")
@auth.login_required
def history():
    rows = db.query(
        "SELECT c.*, u.full_name AS who, p.name AS project"
        " FROM calculations c LEFT JOIN users u ON u.id=c.run_by"
        " LEFT JOIN projects p ON p.id=c.project_id ORDER BY c.run_at DESC LIMIT 200")
    return render_template("history.html", rows=[dict(r) for r in rows])


@app.route("/calculation/<int:cid>")
@auth.login_required
def calculation_detail(cid):
    row = db.query("SELECT c.*, u.full_name AS who FROM calculations c"
                   " LEFT JOIN users u ON u.id=c.run_by WHERE c.id=?", (cid,), one=True)
    if not row:
        abort(404)
    data = dict(row)
    data["results"] = json.loads(data["results"])
    data["inputs"] = json.loads(data["inputs"])
    return render_template("detail.html", calc=data)


@app.route("/calculation/<int:cid>/print")
@auth.login_required
def calculation_print(cid):
    """گزارش چاپی A4 — با Ctrl+P و انتخاب «Save as PDF» فایل PDF می‌دهد."""
    row = db.query("SELECT c.*, u.full_name AS who, u.initials FROM calculations c"
                   " LEFT JOIN users u ON u.id=c.run_by WHERE c.id=?", (cid,), one=True)
    if not row:
        abort(404)
    data = dict(row)
    data["results"] = json.loads(data["results"])
    data["inputs"] = json.loads(data["inputs"])
    from structural import validation
    return render_template("print.html", calc=data, conf=st.load(), options=design_options(),
                           validation=validation.run() if data["results"].get("structures") else None)


# ================================================================= ردیف تجهیزات
@app.route("/bay")
@auth.login_required
def bay_page():
    subs = db.query("SELECT s.*, p.name AS project FROM substations s"
                    " LEFT JOIN projects p ON p.id=s.project_id ORDER BY s.code")
    return render_template("bay.html", defaults=asdict(_office_config()),
                           substations=[dict(r) for r in subs], voltages=VOLTAGE_LEVELS,
                           guide=__import__("guide").by_id())


@app.post("/api/bay")
@auth.requires("engineer")
def api_bay():
    """
    طراحی همه فونداسیون‌های یک ردیف با محدودیت فضا (bay.py)؛ هر پی به‌صورت یک محاسبه
    کامل ثبت می‌شود (گزارش و نقشه جدا) و کی‌پلن ردیف به DXF ساخته می‌شود.
    """
    import bay as B
    form = request.get_json(silent=True)
    if not isinstance(form, dict):
        return jsonify({"error": "درخواست باید JSON باشد."}), 400
    try:
        cfg = config_from_form(form)
        voltage = form.get("voltage") or "63"
        if voltage not in VOLTAGE_LEVELS:
            raise InputError("سطح ولتاژ نامعتبر است")
        gap = _number(form, "gap") or 0.20
        span = _number(form, "merge_span")
        span = 1.6 if span is None else span
        L_hi = _number(form, "L_max") or 6.0
        if not (0 <= gap <= 2 and 0 <= span <= 5 and 1 <= L_hi <= 12):
            raise InputError("فاصله آزاد، حد ادغام یا حداکثر طول پی خارج از محدوده است")
        result = B.design_bay(form.get("chain") or "", cfg, voltage, gap, L_hi=L_hi,
                              merge_span=span)
        B.verify(result, cfg)
    except (TypeError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 400

    units = []
    for u in result["units"]:
        Bw, L, tf, vol = u.choice
        res, seis, des, qty, bbs = u.out
        eq = u.layout.main
        payload = to_dict(res, seis, des, qty, bbs, eq, u.cfg)
        cid = _store(form, eq, u.cfg, B.unit_spec(u), payload)
        units.append({"label": u.label, "centre": round(u.centre, 3), "L": L, "B": Bw, "tf": tf,
                      "volume": round(vol, 3), "ok": payload["ok"], "id": cid,
                      "pedestals": [[round(px, 3), round(py, 3)] for g in u.layout.groups
                                    for px, py in g.positions],
                      "stations": [s.label for s in u.stations],
                      "fs": round(res.checks[0].value, 2), "q": round(res.checks[1].value, 2),
                      "steel": [{"tag": res.layout.groups[gi].eq.tag, "chord": d.sections["chord"],
                                 "ratio": round(float(d.max_ratio), 2)}
                                for gi, d in res.structures]})
    name = f"bay_{secrets.token_hex(4)}.dxf"
    try:
        B.plan_dxf(result, str(OUT / name))
        plan = {"name": name, "url": url_for("download", name=name)}
    except ImportError:
        plan = None
    auth.record("طراحی ردیف تجهیزات", "bay", None, {"units": len(units)})
    return jsonify({"units": units, "b": cfg.foundation.b, "gap": result["gap"],
                    "stations": [{"label": s.label, "pos": round(s.pos, 3), "width": s.width,
                                  "equipment": bool(s.keys)} for s in result["stations"]],
                    "merges": [list(m) for m in result["merges"]], "plan": plan,
                    "volume": round(sum(u["volume"] for u in units), 2)})


# ================================================================= راهنما
@app.route("/guide")
@auth.login_required
def guide_page():
    from guide import GUIDE
    return render_template("guide.html", guide=GUIDE)


# ================================================================= روش تحلیل سازه
@app.route("/method")
@auth.login_required
def method_page():
    from structural import validation
    from structural.loads import COMBO_TITLES
    return render_template("method.html", v=validation.run(), combos=COMBO_TITLES)


# ================================================================= مقایسه با SAP
@app.route("/verify")
@auth.login_required
def verify_page():
    return render_template("verify.html")


@app.post("/api/verify")
@auth.requires("engineer")
def api_verify():
    """
    مدل SAP خود دفتر با حل‌کننده برنامه تحلیل و طراحی می‌شود؛ اگر خروجی SAP همان مدل هم
    داده شود، هر عکس‌العمل، تغییرمکان و نسبت تنش کنار عدد SAP قرار می‌گیرد.
    """
    import shutil
    from structural import sapcheck
    fm, fr = request.files.get("model"), request.files.get("results")
    if not fm or not fm.filename.lower().endswith((".s2k", ".$2k", ".txt")):
        return jsonify({"error": "مدل SAP را با قالب متنی بدهید: در SAP از File ← Export ← "
                                 "SAP2000 .s2k Text File (یا Save As با پسوند $2k)."}), 400
    if fr and fr.filename and not fr.filename.lower().endswith((".xlsx", ".xlsm", ".s2k", ".$2k", ".txt")):
        return jsonify({"error": "خروجی SAP باید Excel (xlsx) یا متن (s2k/txt) باشد."}), 400
    folder = OUT / f"verify_{secrets.token_hex(8)}"
    folder.mkdir()
    try:
        mp = folder / ("model" + Path(fm.filename).suffix.lower())
        fm.save(mp)
        rp = None
        if fr and fr.filename:
            rp = folder / ("results" + Path(fr.filename).suffix.lower())
            fr.save(rp)
        report = sapcheck.compare(str(mp), str(rp) if rp else None)
        report["model"] = Path(fm.filename).stem
        name = f"program_results_{Path(fm.filename).stem}_{secrets.token_hex(3)}.xlsx"
        sapcheck.write_excel(report["program"], str(OUT / name), Path(fm.filename).stem)
    except sapcheck.NotResults as exc:
        return jsonify({"error": str(exc)}), 400
    except (KeyError, ValueError) as exc:
        return jsonify({"error": f"فایل خوانده نشد: {exc}"}), 400
    except Exception as exc:
        app.logger.exception("verify")
        return jsonify({"error": f"فایل خوانده نشد ({type(exc).__name__}: {exc})"}), 500
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    auth.record("مقایسه با SAP", "verify", None, {"model": fm.filename,
                                                  "results": fr.filename if fr else None,
                                                  "ok": report["ok"]})
    prog = report.pop("program")
    report["program_counts"] = {k: len(v) for k, v in prog.items()}
    report["program_steel"] = sorted(prog["steel"], key=lambda r: -r["Ratio"])[:12]
    # نمونه برای دیدن با چشم: بزرگ‌ترین نیروهای قائم تکیه‌گاه در ترکیب‌های طراحی
    design = set(report["design_combos"])
    big = [r for r in prog["reactions"] if r["OutputCase"] in design] or prog["reactions"]
    report["program_reactions"] = sorted(big, key=lambda r: -abs(r["F3"]))[:8]
    if "reactions" in report:
        rows = [r for r in report["reactions"] if r["case"] in design] or report["reactions"]
        report["sample_reactions"] = sorted(rows, key=lambda r: -abs(r["sap"]["F3"] or 0))[:8]
    for k in ("reactions", "displacements"):
        if k in report:
            report[k + "_count"] = len(report[k])
            report[k] = sorted(report[k], key=lambda r: -r["error"])[:60]
    if "steel" in report:
        report["steel_count"] = len(report["steel"])
    report["excel"] = url_for("download", name=name)
    return jsonify(report)


# ================================================================= آیین‌نامه
@app.route("/codes")
@auth.login_required
def codes():
    profiles = cp.listing(active_only=False)
    a = request.args.get("a", type=int)
    b = request.args.get("b", type=int)
    changes = cp.diff(a, b) if a and b and a != b else None
    return render_template("codes.html", profiles=profiles,
                           changes=changes, a=a, b=b)


# ================================================================= کاتالوگ
@app.route("/catalog")
@auth.login_required
def catalog():
    rows = db.query("SELECT e.*, u.full_name AS who FROM equipment_catalog e"
                    " LEFT JOIN users u ON u.id=e.created_by ORDER BY e.tag")
    items = []
    for r in rows:
        d = dict(r)
        d["data"] = json.loads(d["data"])
        items.append(d)
    return render_template("catalog.html", items=items,
                           builtin={k: asdict(v) for k, v in ALL_EQUIPMENT.items()})


@app.post("/catalog/import-builtin")
@auth.requires("engineer")
def catalog_import():
    uid = auth.current_user()["id"]
    added = 0
    for tag, eq in ALL_EQUIPMENT.items():
        exists = db.query("SELECT id FROM equipment_catalog WHERE tag=? AND"
                          " IFNULL(manufacturer,'')='' AND IFNULL(model,'')=''",
                          (tag,), one=True)
        if exists:
            continue
        db.execute("INSERT INTO equipment_catalog (tag, title, data, source,"
                   " created_by, created_at) VALUES (?,?,?,?,?,?)",
                   (tag, eq.title, json.dumps(asdict(eq), ensure_ascii=False),
                    eq.source, uid, db.now()))
        added += 1
    auth.record("درج کاتالوگ پایه", "equipment_catalog", None, {"added": added})
    flash(f"{added} تجهیز از کاتالوگ پایه افزوده شد.", "ok")
    return redirect(url_for("catalog"))


# ================================================================= حساب کاربری
@app.route("/account", methods=["GET", "POST"])
@auth.login_required
def account():
    if request.method == "POST":
        f = request.form
        if f.get("new", "") != f.get("confirm", ""):
            flash("رمز تازه و تکرارش یکی نیستند.", "error")
        else:
            try:
                auth.change_password(auth.current_user()["id"], f.get("old", ""), f.get("new", ""))
                auth.record("تغییر رمز عبور", "users", auth.current_user()["id"])
                flash("رمز عبور عوض شد.", "ok")
                return redirect(url_for("dashboard"))
            except ValueError as exc:
                flash(str(exc), "error")
    return render_template("account.html", min_len=auth.MIN_PASSWORD)


# ================================================================= کاربران
@app.route("/users")
@auth.requires("admin")
def users():
    rows = db.query("SELECT * FROM users ORDER BY id")
    return render_template("users.html", rows=[dict(r) for r in rows])


@app.post("/users/new")
@auth.requires("admin")
def users_new():
    f = request.form
    try:
        uid = auth.create_user(f["username"], f["full_name"], f["password"],
                               f.get("role", "engineer"), f.get("initials"))
        auth.record("ایجاد کاربر", "users", uid, {"username": f["username"]})
        flash("کاربر ساخته شد.", "ok")
    except (ValueError, KeyError) as exc:
        flash(f"کاربر ساخته نشد: {exc}", "error")
    except Exception as exc:
        flash("کاربر ساخته نشد — نام کاربری تکراری است؟" if "UNIQUE" in str(exc)
              else f"کاربر ساخته نشد: {exc}", "error")
    return redirect(url_for("users"))


@app.route("/audit")
@auth.requires("admin")
def audit():
    rows = db.query("SELECT a.*, u.full_name AS who FROM audit_log a"
                    " LEFT JOIN users u ON u.id=a.user_id ORDER BY a.at DESC LIMIT 300")
    return render_template("audit.html", rows=[dict(r) for r in rows])


@app.errorhandler(400)
def bad_request(exc):
    return render_template("error.html", code=400,
                           message=getattr(exc, "description", "درخواست نامعتبر")), 400


@app.errorhandler(403)
def forbidden(_):
    return render_template("error.html", code=403,
                           message="این بخش برای نقش کاربری شما باز نیست."), 403


@app.errorhandler(404)
def missing(_):
    return render_template("error.html", code=404,
                           message="چنین صفحه‌ای وجود ندارد."), 404


def bootstrap():
    db.init_db()
    creds = auth.seed_admin()
    admin = db.query("SELECT id FROM users WHERE role='admin' ORDER BY id LIMIT 1", one=True)
    cp.seed_defaults(admin["id"] if admin else None)
    return creds


if __name__ == "__main__":
    creds = bootstrap()
    if creds:
        print(f"کاربر مدیر ساخته شد — {creds}  (رمز را بعد از اولین ورود عوض کنید)")
    port = int(os.environ.get("FOUNDATION_PORT", 5000))
    url = f"http://127.0.0.1:{port}"
    print(f"سامانه روی {url} بالا آمد — برای بستن این پنجره را ببندید")
    if os.environ.get("FOUNDATION_OPEN_BROWSER") == "1":
        import threading, webbrowser
        threading.Timer(1.5, lambda: webbrowser.open(url)).start()
    app.run(debug=os.environ.get("FOUNDATION_DEBUG") == "1", port=port)
