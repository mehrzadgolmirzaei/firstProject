"""
مقایسه عددی با SAP2000 — اثبات مستقل درستی حل‌کننده و طراحی برنامه.

روند:
    ۱. مدل SAP خود دفتر (.s2k یا .$2k) با حل‌کننده برنامه تحلیل و با AISC-ASD89 کنترل می‌شود.
    ۲. خروجی همان مدل از خود SAP (Excel یا متن، از Display ← Show Tables ← Export) خوانده می‌شود:
       Joint Reactions، Joint Displacements، Steel Design 1 - Summary Data - AISC-ASD89.
    ۳. هر عدد برنامه کنار عدد SAP؛ اختلاف هر ردیف نسبت به بزرگ‌ترین مقدار همان حالت بار.

خروجی برنامه به همان قالب جدول‌های SAP هم نوشته می‌شود (Excel) تا مقایسه دستی هم ممکن باشد.
"""
import math
from collections import defaultdict
from pathlib import Path

from . import asd89
from .frame import analyse
from .s2k import read_model, read_tables

REACTION = ("F1", "F2", "F3", "M1", "M2", "M3")
DISPLACEMENT = ("U1", "U2", "U3", "R1", "R2", "R3")
TABLES = {"reactions": "Joint Reactions", "displacements": "Joint Displacements",
          "steel": "Steel Design 1 - Summary Data - AISC-ASD89"}
TOL = 0.01                   # ۱٪ بزرگ‌ترین مقدار همان حالت بار (اعداد SAP گرد‌شده چاپ می‌شوند)
RATIO_TOL = 0.01             # اختلاف مطلق نسبت تنش


def _key(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ------------------------------------------------------------------ برنامه
def run(model_path):
    """مدل SAP ← (مدل، نتایج تحلیل، کنترل‌های ASD89، ترکیب‌های طراحی)."""
    model = read_model(model_path)
    res = analyse(model)
    combos = [c for c in (getattr(model, "design_combos", None) or model.combos) if c in res.forces]
    checks = asd89.design(model, res, combos=combos) if combos else {}
    return model, res, checks, combos


def _case_name(c):
    return c[2:] if c.startswith("P:") else c


def program_tables(model, res, checks, combos):
    """نتایج برنامه به قالب جدول‌های SAP."""
    cases = [c for c in res.cases() if not (c.startswith("P:") and c[2:] in res.cases())]
    reactions, disps = [], []
    for c in cases:
        for n, r in res.reactions[c].items():
            reactions.append({"Joint": n, "OutputCase": _case_name(c),
                              **{k: float(v) for k, v in zip(REACTION, r)}})
        for n in model.nodes:
            d = res.displacement(c, n)
            disps.append({"Joint": n, "OutputCase": _case_name(c),
                          **{k: float(v) for k, v in zip(DISPLACEMENT, d)}})
    steel = []
    for m in model.members:
        ch = checks.get(m.name)
        if ch is None:
            continue
        kind = "PMM" if ch.governing == ch.ratio else ("Shear" if ch.governing == ch.shear_ratio
                                                        else "KL/r")
        steel.append({"Frame": m.name, "DesignSect": m.section.name, "Ratio": float(ch.governing),
                      "RatioType": kind, "Combo": ch.combo, "Location": float(ch.station),
                      "Equation": ch.equation})
    return {"reactions": reactions, "displacements": disps, "steel": steel}


def write_excel(tables, path, title=""):
    """جدول‌ها به Excel با همان سرآیند SAP (TABLE: …، ستون‌ها، واحدها)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    units = {"F": "Kgf", "M": "Kgf-m", "U": "m", "R": "Radians", "Location": "m"}
    wb = Workbook()
    wb.remove(wb.active)
    for key, rows in tables.items():
        ws = wb.create_sheet(TABLES[key][:31])
        ws.append([f"TABLE:  {TABLES[key]}" + (f"  —  {title}" if title else "")])
        ws["A1"].font = Font(bold=True)
        if not rows:
            continue
        cols = list(rows[0])
        ws.append(cols)
        ws.append([units.get(c, units.get(c[0], "")) if c not in ("Joint", "Frame", "OutputCase",
                   "DesignSect", "RatioType", "Combo", "Equation") else "" for c in cols])
        for r in rows:
            ws.append([r[c] for c in cols])
    wb.save(path)
    return path


# ------------------------------------------------------------------ خروجی SAP
def _rows_from_sheet(ws):
    """یک برگه Excel خروجی SAP: ردیف عنوان TABLE، ردیف ستون‌ها، ردیف واحدها، داده."""
    vals = [list(r) for r in ws.iter_rows(values_only=True)]
    title, head, out = None, None, []
    for row in vals:
        cells = ["" if v is None else v for v in row]
        if not any(str(c).strip() for c in cells):
            continue
        first = str(cells[0]).strip()
        if first.upper().startswith("TABLE:"):
            title = first.split(":", 1)[1].strip()
            head = None
            continue
        if head is None:
            head = [str(c).strip() for c in cells]
            continue
        rec = {h: c for h, c in zip(head, cells) if h}
        if all(not isinstance(v, (int, float)) for k, v in rec.items()
               if k in REACTION + DISPLACEMENT + ("Ratio",)):
            if any(str(v).strip() in ("Kgf", "Kgf-m", "m", "Radians", "Tonf", "KN", "N", "mm")
                   for v in rec.values()):
                units = {k: str(v).strip() for k, v in rec.items()}
                if any(u in ("Tonf", "KN", "N", "mm", "Tonf-m", "KN-m", "N-mm")
                       for u in units.values()):
                    raise ValueError("واحد خروجی SAP باید Kgf و m باشد (در SAP واحد را "
                                     "Kgf, m, C کنید و دوباره خروجی بگیرید)")
                continue
        out.append(rec)
    if title is None:
        title = ws.title
    return title, out


def read_sap_results(path):
    """خروجی SAP (xlsx، یا متن .s2k/.$2k/.txt) ← {"reactions"|"displacements"|"steel": [ردیف‌ها]}."""
    path = Path(path)
    raw = defaultdict(list)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            title, rows = _rows_from_sheet(ws)
            raw[title] += rows
    else:
        for title, rows in read_tables(path).items():
            raw[title] += rows
    out = {}
    for key, name in TABLES.items():
        for title, rows in raw.items():
            if title.strip().lower().startswith(name.lower()[:28]):
                out[key] = rows
    return out


# ------------------------------------------------------------------ مقایسه
def _num(v):
    try:
        return float(str(v).replace("/", ".")) if not isinstance(v, (int, float)) else float(v)
    except ValueError:
        return None


def _compare_joint_table(ours, sap, cols):
    """ردیف‌به‌ردیف با کلید (گره، حالت)؛ اختلاف نسبت به بزرگ‌ترین مقدار آن حالت و آن نوع (نیرو/لنگر)."""
    mine = {(_key(r["Joint"]), _key(r["OutputCase"])): r for r in ours}
    scale = defaultdict(float)
    rows = []
    for r in sap:
        if str(r.get("StepType", "") or "").strip() not in ("", "None"):
            continue                                   # پوش (Max/Min) قابل مقایسه تک‌به‌تک نیست
        k = (_key(r.get("Joint")), _key(r.get("OutputCase")))
        if k not in mine:
            continue
        a = {c: _num(r.get(c)) for c in cols}
        rows.append((k, a, mine[k]))
        for c in cols:
            grp = (k[1], c[0])
            for v in (a[c], mine[k][c]):
                if v is not None:
                    scale[grp] = max(scale[grp], abs(v))
    out, worst = [], 0.0
    for k, a, b in rows:
        diffs = {}
        for c in cols:
            if a[c] is None:
                continue
            s = scale[(k[1], c[0])]
            diffs[c] = abs(a[c] - b[c]) / s if s > 1e-9 else 0.0
        e = max(diffs.values(), default=0.0)
        worst = max(worst, e)
        out.append({"joint": k[0], "case": k[1], "sap": a, "program": {c: b[c] for c in cols},
                    "error": float(e), "ok": bool(e <= TOL)})
    return out, worst


class NotResults(ValueError):
    """فایل خروجی داده‌شده خودش یک مدل است، نه نتایج تحلیل SAP."""


def compare(model_path, sap_path):
    model, res, checks, combos = run(model_path)
    ours = program_tables(model, res, checks, combos)
    if sap_path:
        sap = read_sap_results(sap_path)
        if not sap and Path(sap_path).suffix.lower() not in (".xlsx", ".xlsm") \
                and "JOINT COORDINATES" in read_tables(sap_path):
            raise NotResults("فایل دوم یک «مدل» SAP است، نه «خروجی تحلیل». فایل دوم باید جدول‌های "
                             "نتیجه باشد که بعد از Run Analysis و Steel Design از SAP گرفته می‌شود "
                             "(Joint Reactions و Steel Design Summary).")
    else:
        sap = read_sap_results(model_path)      # مدلی که جدول‌های نتیجه را هم همراه دارد
    report = {"model": model.name, "joints": len(model.nodes), "members": len(model.members),
              "cases": len(model.patterns), "combos": len(model.combos),
              "design_combos": combos, "program": ours, "found": sorted(sap)}
    if "reactions" in sap:
        report["reactions"], report["reactions_err"] = _compare_joint_table(
            ours["reactions"], sap["reactions"], REACTION)
    if "displacements" in sap:
        report["displacements"], report["displacements_err"] = _compare_joint_table(
            ours["displacements"], sap["displacements"], DISPLACEMENT)
    if "steel" in sap:
        mine = {_key(r["Frame"]): r for r in ours["steel"]}
        rows, worst = [], 0.0
        for r in sap["steel"]:
            f = _key(r.get("Frame"))
            v = _num(r.get("Ratio"))
            if f not in mine or v is None:
                continue
            d = float(abs(v - mine[f]["Ratio"]))
            worst = max(worst, d)
            rows.append({"frame": f, "section": r.get("DesignSect", ""), "sap": v,
                         "sap_combo": _key(r.get("Combo", "")), "sap_type": r.get("RatioType", ""),
                         "program": mine[f]["Ratio"], "program_combo": mine[f]["Combo"],
                         "program_type": mine[f]["RatioType"], "diff": d, "ok": bool(d <= RATIO_TOL)})
        rows.sort(key=lambda x: -x["diff"])
        report["steel"], report["steel_err"] = rows, worst
    counted = [report.get(k) for k in ("reactions", "displacements", "steel") if report.get(k)]
    report["compared"] = sum(len(x) for x in counted)
    report["ok"] = bool(counted) and all(r["ok"] for x in counted for r in x)
    return report
